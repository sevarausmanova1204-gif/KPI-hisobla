"""Excel (.xlsx) bilan ishlash: offlayn sinov uchun o'qish va maosh eksporti."""
from __future__ import annotations

from datetime import date

from . import config as C
from .calc import SalaryRow
from .formulas import MONTHLY_HEADERS
from .models import HEADERS, Dataset, parse_dataset, parse_settings


def load_xlsx(path: str) -> tuple[Dataset, C.Settings]:
    """Sheets bilan bir xil varaq nomlari va ustunlari bo'lgan .xlsx faylni o'qiydi."""
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    data = {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
    return parse_dataset(data), parse_settings(data.get(C.SH_SETTINGS))


def write_template(path: str) -> None:
    """Bo'sh shablon: sarlavhalar va standart Sozlamalar."""
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    for title, headers in HEADERS.items():
        ws = wb.create_sheet(title)
        ws.append(headers)
        if title == C.SH_SETTINGS:
            for key, value, note in C.DEFAULT_SETTINGS:
                if isinstance(value, str) and value.startswith("="):
                    value = None
                ws.append([key, value, note])
    wb.save(path)


def export_salary(path: str, month: date, rows: list[SalaryRow]) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    from .sheets import salary_cells
    wb = Workbook()
    ws = wb.active
    ws.title = f"Maosh {month.strftime('%Y-%m')}"
    ws.append([f"Oylik maosh — {month.strftime('%m.%Y')}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(MONTHLY_HEADERS)
    for cell in ws[3]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append(salary_cells(r))
    ws.append([])
    ws.append(["Umumiy fond"] + [""] * (len(MONTHLY_HEADERS) - 2) + [sum(r.total for r in rows)])
    for row in ws.iter_rows(min_row=4):
        for cell in row:
            if isinstance(cell.value, (int, float)) and abs(cell.value) >= 1000:
                cell.number_format = "#,##0"
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 14
    ws.column_dimensions["A"].width = 18
    wb.save(path)
    return path
